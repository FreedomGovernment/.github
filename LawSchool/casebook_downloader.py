#!/usr/bin/env python3
"""
OpenCasebook H2O Casebook Downloader & Converter

Downloads featured casebooks from opencasebook.org and converts them to Markdown.

Usage:
    python3 casebook_downloader.py              # Download all featured casebooks
    python3 casebook_downloader.py --id 109     # Download a specific casebook
    python3 casebook_downloader.py --dry-run    # Show what would be downloaded

The site is behind Cloudflare. This script tries multiple strategies:
1. Direct download (works if you're not behind Cloudflare)
2. Wayback Machine archive (works for archived pages)
3. Manual: download DOCX via browser, then convert with --convert-docx

Requires: python3, beautifulsoup4, lxml, python-docx
    pip3 install beautifulsoup4 lxml python-docx
"""

import argparse
import json
import os
import re
import subprocess
import sys
import time
from pathlib import Path
from urllib.request import Request, urlopen
from urllib.error import URLError, HTTPError

# Try imports
try:
    from bs4 import BeautifulSoup, NavigableString
except ImportError:
    print("ERROR: beautifulsoup4 not installed. Run: pip3 install beautifulsoup4")
    sys.exit(1)

try:
    import docx
except ImportError:
    docx = None

# Category mapping (from opencasebook.org/featured/)
CATEGORIES = {
    '1L-Courses': {
        'Civil-Procedure': {
            '5196': 'Civil Procedure (Abramowicz)',
            '1245': 'Civil Procedure (Cohen)',
            '4444': 'U.S. Civil Procedure for International Students (Campbell)',
        },
        'Constitutional-Law': {
            '5182': 'Constitutional Law Precedents (Coble)',
            '531': 'Constitutional Law (Lessig)',
            '6029': 'Constitutional Structures (Gomez-Velez)',
        },
        'Contracts': {
            '4201': 'Contracts (Coates)',
            '11568': 'Contracts Casebook (Bar-Gill)',
            '3665': 'Contracts: Happiness and Heartbreak (Kennedy)',
        },
        'Criminal-Law': {
            '3553': 'Crim Law (Ball & Oberman)',
            '5587': 'Criminal Law (Brown)',
            '6083': 'Federal Criminal Law (Brown)',
            '5713': 'Criminal Advocacy (Hochman Bloom)',
            '2459': 'Criminal Law (Suk Gersen)',
        },
        'Property': {
            '2202': 'Property Law (White & Bratspies)',
        },
        'Torts': {
            '379': 'Torts (Karlan)',
            '2566': 'Torts! (Zittrain & Weinstock)',
            '187': 'Torts (Ziegler)',
        },
    },
    'Upper-Level-Courses': {
        'Corporations': {
            '237': 'An Introduction to the Law of Corporations (Quinn)',
            '261': 'Corporations (Spamann & Frankenreiter)',
        },
        'Employment-Law': {
            '3728': 'Employment Discrimination (Buzuvis)',
            '5444': 'Employment Law (Ziegler)',
        },
        'Evidence': {
            '5959': 'Advanced Evidence (Chernoff)',
            '6094': 'Evidence (Chernoff)',
            '230': 'Federal Rules of Evidence (Coble)',
        },
        'Lawyering': {
            '762': 'Lawyering for Justice (Bavitz)',
        },
        'Technology-and-the-Law': {
            '2554': 'Governing Digital Technology (Zittrain)',
            '4395': 'Music and Digital Media (Bavitz)',
            '6714': 'Technology, Human Rights, and Cybersecurity (Krishnamurthy)',
        },
        'Other-Topics': {
            '3648': 'Advanced Legislation (Rabb)',
            '4746': 'Asian Americans and U.S. Law (Tanaka)',
            '2806': 'Classics in International Trade Law (Steinberg)',
            '588': 'Family Law (Suk Gersen)',
            '5386': 'Gender, Sexuality, and the Law (Ziegler)',
            '5485': 'Immigration Law (Gundavaram)',
            '109': 'Roe at 50 (Ziegler)',
            '29': 'The Role of State Attorneys General (Tierney)',
        },
    },
    'Graduate-and-Undergraduate-Courses': {
        'General': {
            '2175': 'Healthcare Law and Ethics (Coilfi)',
            '882': 'Music IP & Stakeholders (Peters)',
            '960': 'Strategy & Organization in Music (Peters)',
            '4082': 'Criminal Procedure (Field)',
            '4477': 'Law & Society (Field)',
        },
    },
}

USER_AGENT = 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
BASE_DIR = Path(__file__).parent


def slugify(text):
    """Convert a title to a filename-safe slug."""
    text = re.sub(r'[^a-zA-Z0-9\s\-]', '', text)
    text = re.sub(r'\s+', '-', text.strip())
    return text.lower()[:60]


def fetch_url(url, timeout=120, retries=3):
    """Fetch a URL and return the content as bytes."""
    for attempt in range(retries):
        try:
            req = Request(url, headers={'User-Agent': USER_AGENT})
            with urlopen(req, timeout=timeout) as response:
                return response.read()
        except (HTTPError, URLError, TimeoutError, OSError) as e:
            if attempt < retries - 1:
                wait = 5 * (attempt + 1)
                print(f"    Retry {attempt+1}/{retries-1} after error: {e} (waiting {wait}s)")
                time.sleep(wait)
            else:
                raise


def try_direct_download(casebook_id, title):
    """Try to download the casebook directly from opencasebook.org."""
    # Try the printable HTML endpoint
    html_url = f'https://opencasebook.org/casebooks/{casebook_id}/as-printable-html/'
    try:
        content = fetch_url(html_url, timeout=15)
        if len(content) > 5000 and b'casebook' in content.lower():
            return content, 'direct'
    except (HTTPError, URLError) as e:
        print(f"  Direct HTML failed: {e}")
    
    # Try the DOCX export
    docx_url = f'https://opencasebook.org/casebooks/{casebook_id}/export.docx'
    try:
        content = fetch_url(docx_url, timeout=120, retries=3)
        if len(content) > 1000:
            return content, 'docx'
    except (HTTPError, URLError, TimeoutError, OSError) as e:
        print(f"  Direct DOCX failed: {e}")
    
    return None, None


def try_wayback_download(casebook_id):
    """Try to download from the Wayback Machine."""
    html_url = f'https://web.archive.org/web/2025/https://opencasebook.org/casebooks/{casebook_id}/as-printable-html/'
    try:
        content = fetch_url(html_url, timeout=30)
        # Check if it's a real archive (not a 404 page)
        if len(content) > 10000 and b'casebook' in content.lower():
            # Strip Wayback Machine toolbar
            content = strip_wayback_toolbar(content)
            return content, 'wayback'
    except (HTTPError, URLError) as e:
        print(f"  Wayback HTML failed: {e}")
    
    return None, None


def strip_wayback_toolbar(html_bytes):
    """Remove the Wayback Machine toolbar from archived HTML."""
    soup = BeautifulSoup(html_bytes, 'lxml')
    
    # Remove Wayback elements
    for selector in ['#wm-ipp-base', '#wm-ipp', '.wm-ipp', 'script[src*="archive.org"]',
                     'script[src*="wombat"]', 'script[src*="bundle-playback"]',
                     'link[href*="web-static.archive"]', 'link[href*="iconochive"]']:
        for el in soup.select(selector):
            el.decompose()
    
    # Remove the Wayback banner
    banner = soup.find(id='donato')
    if banner:
        banner.decompose()
    
    return str(soup.encode('utf-8'))


def html_to_markdown(html_content):
    """Convert H2O casebook HTML to Markdown."""
    soup = BeautifulSoup(html_content, 'lxml')
    
    # Remove scripts, styles, nav, footer
    for tag in soup(['script', 'style', 'nav', 'footer', 'header', 'noscript']):
        tag.decompose()
    
    # Find the main content
    main = soup.find('main') or soup.find('div', class_='casebook-content') or soup.find('body')
    if not main:
        main = soup
    
    # Build markdown
    lines = []
    
    def process_element(el, depth=0):
        """Recursively process an element into markdown."""
        if isinstance(el, NavigableString):
            text = str(el).strip()
            if text:
                lines.append(text)
            return
        
        if not getattr(el, 'name', None):
            return
        
        tag = el.name
        
        # Headings
        if tag in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
            level = int(tag[1])
            text = el.get_text(strip=True)
            if text:
                lines.append('')
                lines.append(f'{"#" * level} {text}')
                lines.append('')
            return
        
        # Paragraphs
        if tag == 'p':
            text = el.get_text(strip=True)
            if text:
                lines.append('')
                lines.append(text)
                lines.append('')
            return
        
        # Lists
        if tag in ['ul', 'ol']:
            items = el.find_all('li', recursive=False)
            for i, item in enumerate(items):
                text = item.get_text(strip=True)
                prefix = f'{i+1}.' if tag == 'ol' else '-'
                lines.append(f'{prefix} {text}')
            lines.append('')
            return
        
        # Blockquotes
        if tag == 'blockquote':
            text = el.get_text(strip=True)
            if text:
                lines.append('')
                for line in text.split('\n'):
                    lines.append(f'> {line}')
                lines.append('')
            return
        
        # Code blocks
        if tag == 'pre':
            code = el.get_text()
            lines.append('```')
            lines.append(code)
            lines.append('```')
            return
        
        # Tables
        if tag == 'table':
            process_table(el)
            return
        
        # Links
        if tag == 'a':
            href = el.get('href', '')
            text = el.get_text(strip=True)
            if href and text:
                lines.append(f'[{text}]({href})')
            else:
                lines.append(text)
            return
        
        # Strong/emphasis
        if tag == 'strong' or tag == 'b':
            text = el.get_text(strip=True)
            if text:
                lines.append(f'**{text}**')
            return
        
        if tag == 'em' or tag == 'i':
            text = el.get_text(strip=True)
            if text:
                lines.append(f'*{text}*')
            return
        
        # Continue/separator
        if tag in ['hr']:
            lines.append('')
            lines.append('---')
            lines.append('')
            return
        
        # Recurse into children
        for child in el.children:
            process_element(child, depth + 1)
    
    def process_table(table):
        """Convert a table to markdown."""
        rows = table.find_all('tr')
        if not rows:
            return
        
        # Header row
        header_cells = rows[0].find_all(['th', 'td'])
        if header_cells:
            header = [cell.get_text(strip=True) for cell in header_cells]
            lines.append('| ' + ' | '.join(header) + ' |')
            lines.append('|' + '|'.join(['---'] * len(header)) + '|')
        
        # Data rows
        for row in rows[1:]:
            cells = row.find_all(['td', 'th'])
            if cells:
                row_data = [cell.get_text(strip=True) for cell in cells]
                lines.append('| ' + ' | '.join(row_data) + ' |')
        lines.append('')
    
    # Process the main content
    process_element(main)
    
    # Clean up multiple blank lines
    markdown = '\n'.join(lines)
    markdown = re.sub(r'\n{4,}', '\n\n\n', markdown)
    return markdown


def docx_to_markdown(docx_content):
    """Convert a DOCX file to Markdown."""
    if docx is None:
        raise ImportError("python-docx not installed. Run: pip3 install python-docx")
    
    import io
    doc = docx.Document(io.BytesIO(docx_content))
    
    lines = []
    
    for para in doc.paragraphs:
        style = para.style.name if para.style else ''
        text = para.text.strip()
        
        if not text:
            continue
        
        # Headings
        if 'Heading 1' in style or 'Title' in style:
            lines.append(f'# {text}')
        elif 'Heading 2' in style:
            lines.append(f'## {text}')
        elif 'Heading 3' in style:
            lines.append(f'### {text}')
        elif 'Heading 4' in style:
            lines.append(f'#### {text}')
        elif 'Heading 5' in style:
            lines.append(f'##### {text}')
        elif 'Heading 6' in style:
            lines.append(f'###### {text}')
        elif 'List' in style:
            lines.append(f'- {text}')
        else:
            lines.append(text)
        
        lines.append('')
    
    # Process tables
    for table in doc.tables:
        for i, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]
            if i == 0:
                lines.append('| ' + ' | '.join(cells) + ' |')
                lines.append('|' + '|'.join(['---'] * len(cells)) + '|')
            else:
                lines.append('| ' + ' | '.join(cells) + ' |')
        lines.append('')
    
    markdown = '\n'.join(lines)
    markdown = re.sub(r'\n{4,}', '\n\n\n', markdown)
    return markdown


def download_casebook(casebook_id, title, base_dir, dry_run=False):
    """Download and convert a single casebook."""
    # Find the category
    category_path = None
    for major, minors in CATEGORIES.items():
        for minor, books in minors.items():
            if casebook_id in books:
                category_path = base_dir / major / minor
                break
        if category_path:
            break
    
    if not category_path:
        category_path = base_dir / 'Uncategorized'
    
    if dry_run:
        print(f'  [DRY RUN] {casebook_id}: {title} → {category_path}')
        return None
    
    category_path.mkdir(parents=True, exist_ok=True)
    
    slug = slugify(title)
    md_path = category_path / f'{slug}.md'
    
    # Skip if already downloaded
    if md_path.exists() and md_path.stat().st_size > 1000:
        print(f'  ✓ {md_path.name} (already downloaded, {md_path.stat().st_size:,} bytes)')
        return md_path
    
    # Try download strategies
    content, source = try_direct_download(casebook_id, title)
    if content is None:
        content, source = try_wayback_download(casebook_id)
    
    if content is None:
        print(f'  ✗ {title}: Could not download (Cloudflare/Wayback unavailable)')
        return None
    
    # Convert to markdown
    markdown = None
    if source == 'docx':
        print(f'  Converting DOCX to Markdown...')
        try:
            markdown = docx_to_markdown(content)
        except Exception as e:
            print(f'  ✗ DOCX conversion failed: {e}')
            return None
    else:
        print(f'  Converting HTML to Markdown...')
        try:
            markdown = html_to_markdown(content)
        except Exception as e:
            print(f'  ✗ HTML conversion failed: {e}')
            return None
    
    # Write the markdown
    md_path.write_text(markdown, encoding='utf-8')
    
    size = md_path.stat().st_size
    lines = len(markdown.split('\n'))
    print(f'  ✓ {md_path.name} ({size:,} bytes, {lines:,} lines) [source: {source}]')
    
    return md_path


def convert_docx_file(docx_path, output_path=None):
    """Convert a manually-downloaded DOCX file to Markdown."""
    content = Path(docx_path).read_bytes()
    markdown = docx_to_markdown(content)
    
    if output_path is None:
        output_path = Path(docx_path).with_suffix('.md')
    else:
        output_path = Path(output_path)
    
    output_path.write_text(markdown, encoding='utf-8')
    print(f'✓ Converted {docx_path} → {output_path}')
    return output_path


def main():
    parser = argparse.ArgumentParser(description='Download OpenCasebook casebooks as Markdown')
    parser.add_argument('--id', type=str, help='Download a specific casebook by ID')
    parser.add_argument('--dry-run', action='store_true', help='Show what would be downloaded')
    parser.add_argument('--convert-docx', type=str, metavar='FILE',
                        help='Convert a manually-downloaded DOCX file to Markdown')
    parser.add_argument('--output', type=str, metavar='FILE',
                        help='Output path for --convert-docx')
    parser.add_argument('--base-dir', type=str, default=str(BASE_DIR),
                        help=f'Base directory for downloads (default: {BASE_DIR})')
    
    args = parser.parse_args()
    
    # Handle DOCX conversion
    if args.convert_docx:
        convert_docx_file(args.convert_docx, args.output)
        return
    
    base_dir = Path(args.base_dir)
    
    # Get the list of casebooks to download
    if args.id:
        # Find the casebook
        for major, minors in CATEGORIES.items():
            for minor, books in minors.items():
                if args.id in books:
                    download_casebook(args.id, books[args.id], base_dir, args.dry_run)
                    return
        print(f'ERROR: Casebook ID {args.id} not found in featured list')
        sys.exit(1)
    
    # Download all featured casebooks
    total = 0
    success = 0
    failed = 0
    
    print(f'OpenCasebook Featured Casebooks')
    print(f'Base directory: {base_dir}')
    print(f'Dry run: {args.dry_run}')
    print()
    
    for major, minors in CATEGORIES.items():
        print(f'=== {major} ===')
        for minor, books in minors.items():
            print(f'  {minor}:')
            for cid, title in books.items():
                total += 1
                print(f'    {cid}: {title}', end='')
                result = download_casebook(cid, title, base_dir, args.dry_run)
                if result:
                    success += 1
                elif not args.dry_run:
                    failed += 1
                # Be polite to the servers
                if not args.dry_run:
                    time.sleep(1)
        print()
    
    if not args.dry_run:
        print(f'=== Summary: {success}/{total} succeeded, {failed} failed ===')
        if failed > 0:
            print()
            print('For failed downloads, you can:')
            print('  1. Download the DOCX manually from opencasebook.org')
            print('  2. Convert it with: python3 casebook_downloader.py --convert-docx FILE.docx')


if __name__ == '__main__':
    main()
